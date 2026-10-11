#!/usr/bin/env python3
"""Proposition de classeur Excel « PEAnalyze » : mêmes informations que « Suivi des investissements.xlsx »,
réorganisées avec la charte de l'application. Toutes les valeurs dérivées sont des formules.

Usage : python3 scripts/excel-peanalyze.py <appdata.json> <sortie.xlsx>
(appdata.json : export des données de l'app, voir scripts/export-appdata.mts)
"""
import json, sys
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ------------------------------------------------------------------ charte PEAnalyze (tokens.css)
PRIMARY, PETROL, TEXT, NEUTRAL, MODERATE = '007461', '004652', '19212D', '465266', '666E8A'
BG, DIV, DIVIDER = 'F6F7F9', 'EDEEF1', 'AEB2C0'
SUCCESS, SUCCESS_BG, CRIT, CRIT_BG = '137128', 'EAF5EB', 'AE0014', 'FFEDEB'
TINT = 'E6F1EF'  # vert primaire très clair (cartes)
INPUT = '1F3FBF'  # bleu des cellules à saisir
ETF_COLORS = {'ESE': '007461', 'PUST': '0A94A8', 'PAEEM': 'F8D10A', 'ETZ': '004652', 'CACC': 'A0D7A2', 'WPEA': '5159EA'}
EXTRA_COLORS = ['1BCB99', 'BAE3EC', 'C85207', '8E6C8A', 'BCC0CD']
FONT = 'Arial'

EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
EUR_S = '+#,##0.00 "€";-#,##0.00 "€";0.00 "€"'
PRICE = '#,##0.00## "€"'
PCT1, PCT2 = '0.0%', '0.00%'
PCT_S = '+0.0%;-0.0%;0.0%'
PTS = '+0.0" pts";-0.0" pts";0.0" pts"'
DATE = 'dd/mm/yyyy'
MEUR = '#,##0 "M€"'
INT = '#,##0'

thin = Side(style='thin', color=DIV)
B_ROW = Border(bottom=thin)


def font(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def iso(s):
    return datetime.strptime(s[:10], '%Y-%m-%d')


class Sheet:
    def __init__(self, wb, title, heading, subtitle, tab=PRIMARY, width=None):
        self.ws = wb.create_sheet(title)
        ws = self.ws
        ws.sheet_properties.tabColor = tab
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        ncol = width or 14
        for c in range(1, ncol + 1):
            ws.cell(1, c).fill = fill(PETROL)
            ws.cell(2, c).fill = fill(PETROL)
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[2].height = 20
        ws['A1'] = heading
        ws['A1'].font = Font(name=FONT, size=16, bold=True, color='FFFFFF')
        ws['A1'].alignment = Alignment(vertical='center', indent=1)
        ws['A2'] = subtitle
        ws['A2'].font = Font(name=FONT, size=9, color='D5E3E1')
        ws['A2'].alignment = Alignment(vertical='top', indent=1)
        ws.row_dimensions[3].height = 10

    def section(self, cell, text):
        self.ws[cell] = text.upper()
        self.ws[cell].font = Font(name=FONT, size=9, bold=True, color=NEUTRAL)

    def header(self, row, col, labels, height=30):
        for i, lab in enumerate(labels):
            c = self.ws.cell(row, col + i, lab)
            c.font = Font(name=FONT, size=9, bold=True, color='FFFFFF')
            c.fill = fill(PRIMARY)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.ws.row_dimensions[row].height = height

    def put(self, row, col, value, fmt=None, bold=False, color=TEXT, inp=False, align=None, size=10, bg=None, border=True):
        c = self.ws.cell(row, col, value)
        c.font = font(size, bold, INPUT if inp else color)
        if fmt:
            c.number_format = fmt
        if align:
            c.alignment = Alignment(horizontal=align, vertical='center')
        if bg:
            c.fill = fill(bg)
        if border:
            c.border = B_ROW
        return c

    def widths(self, ws_widths):
        for col, w in ws_widths.items():
            self.ws.column_dimensions[col].width = w


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def style_chart(ch, title=None, legend='b'):
    if title:
        ch.title = title
        ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1100, b=True, solidFill=TEXT, latin=DFont(typeface=FONT)))
    if legend is None:
        ch.legend = None
    else:
        ch.legend.position = legend
    ch.style = 2
    tx = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=800, solidFill=MODERATE, latin=DFont(typeface=FONT))), endParaRPr=CharacterProperties(sz=800))])
    for ax in ('x_axis', 'y_axis'):
        a = getattr(ch, ax, None)
        if a is not None:
            a.txPr = tx
            a.delete = False
            a.majorTickMark = 'none'
            a.minorTickMark = 'none'
            a.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=DIV))
    if getattr(ch, 'y_axis', None) is not None and ch.y_axis.majorGridlines is not None:
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=DIV))
    if ch.legend is not None:
        ch.legend.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=800, solidFill=NEUTRAL, latin=DFont(typeface=FONT))), endParaRPr=CharacterProperties(sz=800))])
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def color_series(s, color, line=False, dash=None, width=22000):
    if line:
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = width
        if dash:
            s.graphicalProperties.line.dashStyle = dash
        s.smooth = False
        s.marker = Marker(symbol='none')
    else:
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color


def main(src, dst):
    D = json.load(open(src))
    etfs = D['etfs']
    tickers = [e['id'] for e in etfs]
    colors = {t: ETF_COLORS.get(t, EXTRA_COLORS[i % len(EXTRA_COLORS)]) for i, t in enumerate(tickers)}
    deps = sorted(D['deposits'], key=lambda d: d['date'])
    txs = sorted(D['transactions'], key=lambda t: t['date'])
    vd = D['valuationDate']
    NROWS = 1000  # capacité des tableaux de saisie

    wb = Workbook()
    wb.remove(wb.active)

    # ================================================================ feuilles (ordre des onglets)
    dash = Sheet(wb, 'Tableau de bord', 'PEAnalyze · Tableau de bord', '', width=9)
    pf = Sheet(wb, 'Portefeuille', 'Portefeuille', 'Positions, allocation réelle et cible, parts à acheter, répartitions sectorielle et géographique.', width=16)
    hi = Sheet(wb, 'Historique', 'Historique des opérations', 'Saisissez vos virements et achats dans les cellules en bleu : montants, frais et indicateurs se calculent seuls.', width=14)
    ev = Sheet(wb, 'Évolution', 'Évolution mensuelle', 'Montant épargné et valeur réelle du portefeuille à la fin de chaque mois.', width=13)
    me = Sheet(wb, 'Mes ETF', 'Mes ETF', 'Caractéristiques, stratégie et allocation cible des ETF suivis. Cours : feuille « Cours ».', width=14)
    co = Sheet(wb, 'Cours', 'Cours de clôture', 'Historique quotidien des cours (Euronext Paris) et performance en base 100.', tab=PETROL, width=2 + 2 * len(tickers) + 1)
    cp = Sheet(wb, 'Compositions', 'Compositions des indices', 'Répartition sectorielle et géographique de chaque ETF (poids saisis) et contribution au portefeuille.', tab=PETROL, width=6)
    ca = Sheet(wb, 'Catalogue', 'Catalogue des ETF éligibles au PEA', '', tab=PETROL, width=15)
    si = Sheet(wb, 'Simulateur', "Simulateur d'épargne", "Capitalisation mensuelle, versements par phases, résultat en euros courants et constants.", width=10)

    # ================================================================ Cours
    ws = co.ws
    hist = {}
    for e in etfs:
        for d, p in e['priceHistory']:
            hist.setdefault(d, {})[e['id']] = p
    dates = sorted(hist)
    nC = len(tickers)
    co.header(3, 1, ['Date'] + tickers, height=22)
    base_col = nC + 3
    co.header(3, base_col - 1, ['Base 100'] + tickers, height=22)
    for i, d in enumerate(dates):
        r = 4 + i
        co.put(r, 1, iso(d), DATE)
        for j, t in enumerate(tickers):
            v = hist[d].get(t)
            co.put(r, 2 + j, v, PRICE, inp=True)
            col = L(2 + j)
            co.put(r, base_col + j, f'=IF({col}{r}="","",{col}{r}/INDEX({col}$4:{col}$5000,MATCH(TRUE,INDEX({col}$4:{col}$5000<>"",0),0))*100)', '0.0')
        co.put(r, base_col - 1, f'=A{r}', DATE, color=MODERATE)
    last_c = 4 + len(dates) - 1
    ws.freeze_panes = 'B4'
    co.widths({'A': 12, L(base_col - 1): 12})
    for j in range(nC):
        ws.column_dimensions[L(2 + j)].width = 10
        ws.column_dimensions[L(base_col + j)].width = 9
    ws.column_dimensions[L(nC + 2)].width = 3
    for j, t in enumerate(tickers):
        ws.cell(3, 2 + j).fill = fill(colors[t]) if colors[t] not in ('F8D10A', 'A0D7A2', 'BAE3EC') else fill(PRIMARY)
    name(wb, 'ValDate', f"MAX(Cours!$A$4:$A$5000)")

    # ================================================================ Mes ETF
    ws = me.ws
    me.header(4, 1, ['Ticker', 'Indice', 'Nom', 'ISIN', 'Catégorie', 'Dans la stratégie', 'Allocation cible', 'Frais (TER)', 'Encours', 'Réplication', 'Distribution', 'Dernier cours', 'Date du cours', 'Perf. depuis le début'])
    dv = DataValidation(type='list', formula1='"Oui,Non"', allow_blank=False)
    ws.add_data_validation(dv)
    for i, e in enumerate(etfs):
        r = 5 + i
        t = e['id']
        me.put(r, 1, t, bold=True)
        ws.cell(r, 1).fill = fill(colors[t])
        ws.cell(r, 1).font = font(10, True, 'FFFFFF' if colors[t] not in ('F8D10A', 'A0D7A2', 'BAE3EC') else TEXT)
        me.put(r, 2, e['index'], inp=True)
        me.put(r, 3, e['name'], inp=True)
        me.put(r, 4, e['isin'], inp=True)
        me.put(r, 5, e['category'], inp=True)
        me.put(r, 6, 'Oui' if e['inStrategy'] else 'Non', inp=True, align='center')
        dv.add(ws.cell(r, 6))
        me.put(r, 7, D['targetAllocation'].get(t, 0) if e['inStrategy'] else 0, PCT1, inp=True, align='center')
        me.put(r, 8, e['ter'], PCT2, inp=True)
        me.put(r, 9, e['aumMEur'], MEUR, inp=True)
        me.put(r, 10, e['replication'], inp=True)
        me.put(r, 11, e['distribution'], inp=True)
        colm = f'MATCH($A{r},Cours!$B$3:${L(1 + nC)}$3,0)'
        rng = f'Cours!$B$4:${L(1 + nC)}$5000'
        me.put(r, 12, f'=INDEX({rng},MATCH(9.99E+307,INDEX({rng},0,{colm})),{colm})', PRICE, bold=True)
        me.put(r, 13, f'=INDEX(Cours!$A$4:$A$5000,MATCH(9.99E+307,INDEX({rng},0,{colm})))', DATE)
        me.put(r, 14, f'=L{r}/INDEX({rng},MATCH(TRUE,INDEX(INDEX({rng},0,{colm})<>"",0),0),{colm})-1', PCT_S)
    n_e = len(etfs)
    rt = 5 + n_e
    me.put(rt, 6, 'Total cible', bold=True, align='right', border=False)
    me.put(rt, 7, f'=SUM(G5:G{rt - 1})', PCT1, bold=True, align='center', border=False)
    me.put(rt, 8, f'=IF(ROUND(G{rt},4)=1,"✓ 100 %","⚠ doit faire 100 %")', color=MODERATE, border=False)
    me.widths({'A': 9, 'B': 16, 'C': 44, 'D': 15, 'E': 30, 'F': 11, 'G': 11, 'H': 9, 'I': 11, 'J': 12, 'K': 13, 'L': 12, 'M': 12, 'N': 12})
    ws.freeze_panes = 'B5'
    E_last = 30
    name(wb, 'E_Ticker', f"'Mes ETF'!$A$5:$A${E_last}")
    name(wb, 'E_Indice', f"'Mes ETF'!$B$5:$B${E_last}")
    name(wb, 'E_Strat', f"'Mes ETF'!$F$5:$F${E_last}")
    name(wb, 'E_Cible', f"'Mes ETF'!$G$5:$G${E_last}")
    name(wb, 'E_Cours', f"'Mes ETF'!$L$5:$L${E_last}")
    # graphique base 100
    ch = LineChart()
    style_chart(ch, 'Performance des ETF (base 100 au début de l’historique)')
    for j, t in enumerate(tickers):
        ch.add_data(Reference(co.ws, min_col=base_col + j, min_row=3, max_row=last_c), titles_from_data=True)
        color_series(ch.series[-1], colors[t], line=True, width=15000)
    ch.set_categories(Reference(co.ws, min_col=base_col - 1, min_row=4, max_row=last_c))
    ch.x_axis.number_format = 'mmm yy'
    ch.x_axis.tickLblSkip = 42
    ch.x_axis.tickMarkSkip = 42
    ch.y_axis.number_format = '0'
    ch.width, ch.height = 26, 9
    ws.add_chart(ch, f'B{rt + 3}')

    # ================================================================ Historique
    ws = hi.ws
    hi.section('A4', 'Virements')
    hi.section('D4', 'Achats')
    hi.header(5, 1, ['Date', 'Montant'])
    hi.header(5, 4, ['Date', 'ETF', 'Indice', 'Cours', 'Quantité', 'Montant brut', 'Débit courtier', 'Frais', 'Rembour-sement', 'Frais nets', 'Frais en %'])
    first = 6
    last = first + NROWS - 1
    dv2 = DataValidation(type='list', formula1=f"='Mes ETF'!$A$5:$A${E_last}", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f'E{first}:E{last}')
    for i in range(NROWS):
        r = first + i
        if i < len(deps):
            hi.put(r, 1, iso(deps[i]['date']), DATE, inp=True)
            hi.put(r, 2, deps[i]['amount'], EUR, inp=True)
        else:
            for c, f in ((1, DATE), (2, EUR)):
                ws.cell(r, c).number_format = f
                ws.cell(r, c).font = font(10, color=INPUT)
        if i < len(txs):
            t = txs[i]
            hi.put(r, 4, iso(t['date']), DATE, inp=True)
            hi.put(r, 5, t['etfId'], inp=True, align='center')
            hi.put(r, 7, t['unitPrice'], PRICE, inp=True)
            hi.put(r, 8, t['quantity'], INT, inp=True)
            hi.put(r, 10, t['grossAmount'], EUR, inp=True)
            hi.put(r, 12, t['feeRefund'], EUR, inp=True)
        else:
            for c, f in ((4, DATE), (5, None), (7, PRICE), (8, INT), (10, EUR), (12, EUR)):
                if f:
                    ws.cell(r, c).number_format = f
                ws.cell(r, c).font = font(10, color=INPUT)
        if i < max(len(txs) + 200, 300):
            hi.put(r, 6, f'=IF(E{r}="","",IFERROR(INDEX(E_Indice,MATCH(E{r},E_Ticker,0)),"?"))', color=NEUTRAL, border=i < len(txs))
            hi.put(r, 9, f'=IF(D{r}="","",G{r}*H{r})', EUR, border=i < len(txs))
            hi.put(r, 11, f'=IF(D{r}="","",J{r}-I{r})', EUR, border=i < len(txs))
            hi.put(r, 13, f'=IF(D{r}="","",K{r}-L{r})', EUR, border=i < len(txs))
            hi.put(r, 14, f'=IF(OR(D{r}="",I{r}=0),"",M{r}/I{r})', PCT2, border=i < len(txs))
        # flux pour le TRI (colonnes masquées)
        ws.cell(r, 16, f'=IF(A{r}="",MIN($A${first}:$A${last}),A{r})').number_format = DATE
        ws.cell(r, 17, f'=IF(B{r}="",0,-B{r})').number_format = EUR
    ws.cell(first - 1, 16, f'=MIN($A${first}:$A${last})').number_format = DATE
    ws.cell(first - 1, 17, 0)
    ws.cell(last + 1, 16, '=ValDate').number_format = DATE
    ws.cell(last + 1, 17, '=Valorisation').number_format = EUR
    ws.cell(4, 16, 'Flux pour le TRI').font = font(8, color=MODERATE)
    for c in ('P', 'Q'):
        ws.column_dimensions[c].hidden = True
    hi.widths({'A': 12, 'B': 12, 'C': 3, 'D': 12, 'E': 8, 'F': 14, 'G': 11, 'H': 9, 'I': 13, 'J': 13, 'K': 9, 'L': 11, 'M': 10, 'N': 9, 'O': 3})
    ws.freeze_panes = f'A{first}'
    ws.conditional_formatting.add(f'N{first}:N{last}', CellIsRule(operator='greaterThan', formula=['0.005'], font=Font(name=FONT, color=CRIT)))
    rng = lambda c: f'Historique!${c}${first}:${c}${last}'
    for nm, c in (('V_Date', 'A'), ('V_Montant', 'B'), ('A_Date', 'D'), ('A_ETF', 'E'), ('A_Cours', 'G'), ('A_Qte', 'H'), ('A_Brut', 'I'),
                  ('A_Debit', 'J'), ('A_Frais', 'K'), ('A_Remb', 'L'), ('A_FraisNets', 'M'), ('A_FraisPct', 'N')):
        name(wb, nm, rng(c))
    name(wb, 'TRI_Dates', f'Historique!$P${first - 1}:$P${last + 1}')
    name(wb, 'TRI_Flux', f'Historique!$Q${first - 1}:$Q${last + 1}')
    ws['J5'].comment = Comment('Montant réellement débité par le courtier, frais inclus.', 'PEAnalyze')
    ws['L5'].comment = Comment('Frais remboursés par le courtier (offres promotionnelles) : ils reviennent sur le solde espèces.', 'PEAnalyze')

    # ================================================================ Portefeuille
    ws = pf.ws
    order = [e['id'] for e in etfs]
    pf.header(4, 1, ['Indice', 'Ticker', 'Stratégie', 'Parts', 'Coût de revient', 'PRU', 'Cours', 'Valeur', 'Plus-value', 'Plus-value %', 'Poids réel', 'Cible', 'Écart', 'Valeur cible', 'Parts cibles', 'Parts à acheter'])
    p0 = 5
    for i, t in enumerate(order):
        r = p0 + i
        m = f'MATCH($B{r},E_Ticker,0)'
        pf.put(r, 1, f'=INDEX(E_Indice,{m})', bold=True)
        pf.put(r, 2, t, align='center')
        ws.cell(r, 2).fill = fill(colors[t])
        ws.cell(r, 2).font = font(10, True, 'FFFFFF' if colors[t] not in ('F8D10A', 'A0D7A2', 'BAE3EC') else TEXT)
        pf.put(r, 3, f'=IF(INDEX(E_Strat,{m})="Oui","Oui","Hors stratégie")', color=NEUTRAL, align='center')
        pf.put(r, 4, f'=SUMIF(A_ETF,$B{r},A_Qte)', INT, align='center')
        pf.put(r, 5, f'=SUMIF(A_ETF,$B{r},A_Debit)-SUMIF(A_ETF,$B{r},A_Remb)', EUR)
        pf.put(r, 6, f'=IF(D{r}>0,E{r}/D{r},"")', PRICE)
        pf.put(r, 7, f'=INDEX(E_Cours,{m})', PRICE)
        pf.put(r, 8, f'=D{r}*G{r}', EUR, bold=True)
        pf.put(r, 9, f'=H{r}-E{r}', EUR_S)
        pf.put(r, 10, f'=IF(E{r}>0,I{r}/E{r},"")', PCT_S)
        pf.put(r, 11, f'=IF(TotalTitres>0,H{r}/TotalTitres,0)', PCT1, align='center')
        pf.put(r, 12, f'=IF(C{r}="Oui",INDEX(E_Cible,{m}),0)', PCT1, align='center')
        pf.put(r, 13, f'=(K{r}-L{r})*100', PTS, align='center')
        pf.put(r, 14, f'=L{r}*Valorisation', EUR)
        pf.put(r, 15, f'=IF(G{r}>0,ROUND(N{r}/G{r},0),0)', INT, align='center')
        pf.put(r, 16, f'=MAX(0,O{r}-D{r})', INT, bold=True, align='center', color=PRIMARY)
    pt = p0 + len(order)
    pf.put(pt, 1, 'Total titres', bold=True, bg=TINT)
    for c in range(2, 17):
        pf.put(pt, c, None, bg=TINT)
    for c, f, fmt in ((4, 'SUM', INT), (5, 'SUM', EUR), (8, 'SUM', EUR), (9, 'SUM', EUR_S), (11, 'SUM', PCT1), (12, 'SUM', PCT1), (14, 'SUM', EUR), (16, 'SUM', INT)):
        pf.put(pt, c, f'={f}({L(c)}{p0}:{L(c)}{pt - 1})', fmt, bold=True, bg=TINT, align='center' if fmt in (INT, PCT1) else None)
    pf.put(pt, 10, f'=IF(E{pt}>0,I{pt}/E{pt},"")', PCT_S, bold=True, bg=TINT)
    ws.conditional_formatting.add(f'I{p0}:J{pt}', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, color=CRIT)))
    ws.conditional_formatting.add(f'I{p0}:J{pt}', CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=FONT, color=SUCCESS)))
    ws.conditional_formatting.add(f'M{p0}:M{pt - 1}', CellIsRule(operator='lessThan', formula=['-5'], font=Font(name=FONT, color=CRIT, bold=True), fill=fill(CRIT_BG)))
    ws.conditional_formatting.add(f'M{p0}:M{pt - 1}', CellIsRule(operator='greaterThan', formula=['5'], font=Font(name=FONT, color='8A4B00', bold=True), fill=fill('FFF4D6')))
    name(wb, 'TotalTitres', f'Portefeuille!$H${pt}')
    name(wb, 'P_Ticker', f'Portefeuille!$B${p0}:$B${pt - 1}')
    name(wb, 'P_PoidsReel', f'Portefeuille!$K${p0}:$K${pt - 1}')
    pf.put(pt + 1, 1, "Écart : poids réel − cible, en points. Rouge si l'ETF est sous-pondéré de plus de 5 points, orange s'il est sur-pondéré.", size=8, color=MODERATE, border=False)
    pf.widths({'A': 16, 'B': 9, 'C': 13, 'D': 9, 'E': 13, 'F': 10, 'G': 10, 'H': 12, 'I': 12, 'J': 10, 'K': 9, 'L': 8, 'M': 10, 'N': 12, 'O': 8, 'P': 9})
    ws.freeze_panes = 'C5'

    # ---- Compositions (feuille dédiée) puis répartitions du portefeuille
    ws2 = cp.ws
    cp.header(4, 1, ['ETF', 'Type', 'Libellé', 'Poids dans l’indice', 'Poids dans le portefeuille'])
    r = 5
    for e in etfs:
        for kind, lab in (('sectors', 'Secteur'), ('countries', 'Pays')):
            for it in e[kind]:
                cp.put(r, 1, e['id'], align='center', inp=True)
                cp.put(r, 2, lab, inp=True)
                cp.put(r, 3, it['label'], inp=True)
                cp.put(r, 4, it['weight'], PCT2, inp=True)
                cp.put(r, 5, f'=D{r}*IFERROR(INDEX(P_PoidsReel,MATCH(A{r},P_Ticker,0)),0)', PCT2)
                r += 1
    c_last = 2000
    cp.widths({'A': 8, 'B': 9, 'C': 30, 'D': 12, 'E': 14})
    ws2.freeze_panes = 'A5'
    for nm, c in (('C_ETF', 'A'), ('C_Type', 'B'), ('C_Label', 'C'), ('C_Contrib', 'E')):
        name(wb, nm, f'Compositions!${c}$5:${c}${c_last}')

    def expo_block(col, title, kind, rows):
        # libellé fusionné sur 3 colonnes, poids dans la 4e
        r0 = pt + 4
        vc = col + 3
        pf.section(f'{L(col)}{r0}', title)
        pf.header(r0 + 1, col, ['Libellé', '', '', 'Poids'], height=20)
        ws.merge_cells(start_row=r0 + 1, start_column=col, end_row=r0 + 1, end_column=col + 2)
        for i, x in enumerate(rows):
            rr = r0 + 2 + i
            for k in range(3):
                pf.put(rr, col + k, x['label'] if k == 0 else None)
            ws.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + 2)
            pf.put(rr, vc, f'=SUMIFS(C_Contrib,C_Type,"{kind}",C_Label,{L(col)}{rr})', PCT1)
        end = r0 + 1 + len(rows)
        ws.conditional_formatting.add(f'{L(vc)}{r0 + 2}:{L(vc)}{end}', DataBarRule(start_type='num', start_value=0, end_type='max', color='9FD0C7', showValue=True))
        return r0, end

    sec = D['expected']['sectors']
    cty = D['expected']['countries']
    s0, s1 = expo_block(1, 'Répartition sectorielle du portefeuille', 'Secteur', sec)
    g0, g1 = expo_block(6, 'Répartition géographique du portefeuille', 'Pays', cty)
    for (c0, end, col, title, anchor) in ((s0, min(s1, s0 + 11), 1, 'Secteurs (top 10)', 'K'), (g0, min(g1, g0 + 11), 6, 'Pays (top 10)', 'K')):
        bc = BarChart()
        bc.type = 'bar'
        style_chart(bc, title, legend=None)
        bc.add_data(Reference(ws, min_col=col + 3, min_row=c0 + 2, max_row=end), titles_from_data=False)
        bc.set_categories(Reference(ws, min_col=col, min_row=c0 + 2, max_row=end))
        color_series(bc.series[0], PRIMARY)
        bc.y_axis.number_format = '0%'
        bc.x_axis.scaling.orientation = 'maxMin'
        bc.gapWidth = 60
        bc.width, bc.height = 15, 8
        ws.add_chart(bc, f'{anchor}{c0 if col == 1 else c0 + 18}')

    # ================================================================ Évolution
    ws = ev.ws
    first_m = deps[0]['date'][:7] if deps else vd[:7]
    months = []
    y, mo = int(first_m[:4]), int(first_m[5:7])
    while f'{y:04d}-{mo:02d}' <= vd[:7]:
        months.append(date(y, mo, 1))
        mo += 1
        if mo == 13:
            y, mo = y + 1, 1
    ev.header(4, 1, ['Mois', 'Fin de période', 'Versé dans le mois', 'Montant épargné'] + [f'Valeur {t}' for t in tickers] + ['Espèces', 'Valeur réelle', 'Plus-value'])
    for j, t in enumerate(tickers):
        ws.cell(3, 5 + j, t).font = font(7, color=BG)
    cE, cV, cP = 5 + nC, 6 + nC, 7 + nC
    for i, m in enumerate(months):
        r = 5 + i
        ev.put(r, 1, m, 'mmm yyyy', bold=True)
        ev.put(r, 2, f'=MIN(EOMONTH(A{r},0),ValDate)', DATE, color=MODERATE)
        ev.put(r, 3, f'=SUMIFS(V_Montant,V_Date,">="&A{r},V_Date,"<="&B{r})', EUR)
        ev.put(r, 4, f'=SUMIFS(V_Montant,V_Date,"<="&B{r})', EUR, bold=True)
        for j, t in enumerate(tickers):
            cc = L(2 + j)
            ev.put(r, 5 + j, f'=SUMIFS(A_Qte,A_ETF,{L(5 + j)}$3,A_Date,"<="&$B{r})*IFERROR(INDEX(Cours!{cc}$4:{cc}$5000,MATCH($B{r},Cours!$A$4:$A$5000,1)),0)', EUR, color=NEUTRAL)
        ev.put(r, cE, f'=MAX(0,D{r}-SUMIFS(A_Debit,A_Date,"<="&B{r})+SUMIFS(A_Remb,A_Date,"<="&B{r}))', EUR, color=NEUTRAL)
        ev.put(r, cV, f'=SUM({L(5)}{r}:{L(cE)}{r})', EUR, bold=True, color=PRIMARY)
        ev.put(r, cP, f'={L(cV)}{r}-D{r}', EUR_S)
    m_last = 5 + len(months) - 1
    ws.conditional_formatting.add(f'{L(cP)}5:{L(cP)}{m_last}', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, color=CRIT)))
    ws.conditional_formatting.add(f'{L(cP)}5:{L(cP)}{m_last}', CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=FONT, color=SUCCESS)))
    ev.widths({'A': 11, 'B': 12, 'C': 12, 'D': 13, L(cE): 10, L(cV): 13, L(cP): 12})
    for j in range(nC):
        ws.column_dimensions[L(5 + j)].width = 11
    ws.freeze_panes = 'B5'
    lc = LineChart()
    style_chart(lc, 'Montant épargné et valeur réelle')
    lc.add_data(Reference(ws, min_col=4, min_row=4, max_row=m_last), titles_from_data=True)
    color_series(lc.series[0], MODERATE, line=True, dash='dash', width=19050)
    lc.add_data(Reference(ws, min_col=cV, min_row=4, max_row=m_last), titles_from_data=True)
    color_series(lc.series[1], PRIMARY, line=True, width=28575)
    lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=m_last))
    lc.x_axis.number_format = 'mmm yy'
    lc.y_axis.number_format = '#,##0 "€"'
    lc.width, lc.height = 24, 9
    ws.add_chart(lc, f'A{m_last + 3}')

    # ================================================================ Simulateur
    ws = si.ws
    sim = D['simulation']
    si.section('A4', 'Paramètres')
    rows = [('Capital initial', sim['initialCapital'], EUR0), ('Horizon (années)', sim['horizonYears'], '0'), ('Rendement annuel', sim['annualRate'], PCT1), ('Inflation (euros constants)', 0.02, PCT1)]
    for i, (lab, v, fmt) in enumerate(rows):
        si.put(5 + i, 1, lab, color=NEUTRAL)
        si.put(5 + i, 2, v, fmt, inp=True, bold=True, bg=TINT)
    si.put(9, 1, 'Taux mensuel équivalent', color=NEUTRAL)
    si.put(9, 2, '=(1+B7)^(1/12)-1', '0.000%')
    si.section('D4', 'Versements mensuels par phase')
    si.header(5, 4, ["À partir de l'année", 'Versement mensuel'], height=28)
    phases = sorted(sim['dcaPhases'], key=lambda p: p['fromYear'])
    for i in range(10):
        r = 6 + i
        if i < len(phases):
            si.put(r, 4, phases[i]['fromYear'], '0', inp=True, align='center')
            si.put(r, 5, phases[i]['monthly'], EUR0, inp=True)
        else:
            si.put(r, 4, None, '0', inp=True, align='center')
            si.put(r, 5, None, EUR0, inp=True)
    si.put(16, 4, "Classez les phases par année croissante ; une phase s'applique jusqu'à la suivante.", size=8, color=MODERATE, border=False)
    # résultats
    si.section('G4', "Résultat à l'horizon")
    res = [('Capital final', f'=INDEX(D21:D61,MATCH(B6,A21:A61,0))', EUR0), ('Total versé', f'=INDEX(C21:C61,MATCH(B6,A21:A61,0))', EUR0),
           ('Intérêts composés', '=H5-H6', EUR0), ('Capital final en euros constants', f'=INDEX(F21:F61,MATCH(B6,A21:A61,0))', EUR0)]
    for i, (lab, f, fmt) in enumerate(res):
        si.put(5 + i, 7, lab, color=NEUTRAL)
        si.put(5 + i, 8, f, fmt, bold=True, color=PRIMARY if i == 0 else TEXT, size=12 if i == 0 else 10)
    si.section('A19', 'Projection année par année')
    si.header(20, 1, ['Année', 'Versement mensuel', 'Total versé', 'Capital', 'Intérêts cumulés', 'Capital en euros constants'], height=28)
    for yv in range(0, 41):
        r = 21 + yv
        si.put(r, 1, yv, '0', align='center', bold=True)
        hide = f'A{r}>$B$6'
        if yv == 0:
            si.put(r, 2, '=""', EUR0)
            si.put(r, 3, '=$B$5', EUR0)
            si.put(r, 4, '=$B$5', EUR0, bold=True)
        else:
            si.put(r, 2, f'=IF({hide},"",IFERROR(LOOKUP(A{r},$D$6:$D$15,$E$6:$E$15),0))', EUR0)
            si.put(r, 3, f'=IF({hide},"",C{r - 1}+B{r}*12)', EUR0)
            si.put(r, 4, f'=IF({hide},"",IF($B$9=0,D{r - 1}+B{r}*12,D{r - 1}*(1+$B$9)^12+B{r}*((1+$B$9)^12-1)/$B$9))', EUR0, bold=True)
        si.put(r, 5, f'=IF(D{r}="","",D{r}-C{r})', EUR0, color=SUCCESS)
        si.put(r, 6, f'=IF(D{r}="","",D{r}/(1+$B$8)^A{r})', EUR0, color=NEUTRAL)
    ws.conditional_formatting.add('A21:A61', FormulaRule(formula=['A21>$B$6'], font=Font(name=FONT, color=DIVIDER)))
    si.widths({'A': 28, 'B': 14, 'C': 13, 'D': 16, 'E': 16, 'F': 3, 'G': 30, 'H': 14})
    ws.column_dimensions['F'].width = 16
    hz = int(sim['horizonYears'])
    sc = LineChart()
    style_chart(sc, 'Capital projeté et total versé')
    sc.add_data(Reference(ws, min_col=4, min_row=20, max_row=21 + hz), titles_from_data=True)
    color_series(sc.series[0], PRIMARY, line=True, width=28575)
    sc.add_data(Reference(ws, min_col=3, min_row=20, max_row=21 + hz), titles_from_data=True)
    color_series(sc.series[1], MODERATE, line=True, dash='dash', width=19050)
    sc.add_data(Reference(ws, min_col=6, min_row=20, max_row=21 + hz), titles_from_data=True)
    color_series(sc.series[2], '0A94A8', line=True, dash='sysDot', width=15875)
    sc.set_categories(Reference(ws, min_col=1, min_row=21, max_row=21 + hz))
    sc.y_axis.number_format = '#,##0 "€"'
    sc.width, sc.height = 17, 9
    ws.add_chart(sc, 'G11')

    # ================================================================ Catalogue
    ws = ca.ws
    ws['A2'] = f"{len(D['catalog'])} ETF · source : liste justETF relevée le 05/10/2026 (frais, encours et prix d'une part au 30/09/2026). Filtrez avec les flèches des en-têtes."
    heads = ['Indice', 'Nom', 'ISIN', 'Mnémo', 'Place', 'Catégorie', 'Frais', 'Encours', 'Création', 'Composants', 'Réplication', 'Distribution', 'ESG', "Prix d'une part", 'Remarque']
    ca.header(4, 1, heads)
    for i, c in enumerate(D['catalog']):
        r = 5 + i
        vals = [c['index'], c['name'], c['isin'], c['ticker'] or '—', c['exchange'], c['category'], c['ter'], c['aumMEur'], c['inceptionYear'], c['holdingsCount'],
                c['replication'] or '—', c['distribution'] or '—', 'Oui' if c['esg'] else 'Non', c['unitPrice'], c.get('note', '')]
        fmts = [None, None, None, None, None, None, PCT2, MEUR, '0', INT, None, None, None, PRICE, None]
        for j, (v, f) in enumerate(zip(vals, fmts)):
            cell = ca.put(r, 1 + j, v, f, bold=(j == 0), color=MODERATE if j == 14 else TEXT, size=8 if j == 14 else 10,
                          align='center' if j in (3, 6, 8, 12) else None, bg=BG if i % 2 else None)
    ca_last = 4 + len(D['catalog'])
    ws.auto_filter.ref = f'A4:O{ca_last}'
    ws.freeze_panes = 'B5'
    ca.widths({'A': 30, 'B': 46, 'C': 14, 'D': 8, 'E': 15, 'F': 22, 'G': 8, 'H': 10, 'I': 9, 'J': 11, 'K': 12, 'L': 12, 'M': 6, 'N': 11, 'O': 44})
    ws.conditional_formatting.add(f'G5:G{ca_last}', CellIsRule(operator='lessThanOrEqual', formula=['0.002'], font=Font(name=FONT, color=SUCCESS, bold=True)))
    held = {e['isin'] for e in etfs}
    for i, c in enumerate(D['catalog']):
        if c['isin'] in held:
            ws.cell(5 + i, 1).fill = fill(SUCCESS_BG)

    # ================================================================ Tableau de bord
    ws = dash.ws
    ws['A2'] = 'Valorisé au'
    ws['C2'] = '=ValDate'
    ws['C2'].number_format = DATE
    ws['C2'].font = Font(name=FONT, size=9, bold=True, color='FFFFFF')
    ws['C2'].alignment = Alignment(horizontal='left', vertical='top')
    ws['E2'] = 'cours de clôture Euronext Paris'
    ws['E2'].font = Font(name=FONT, size=9, color='D5E3E1')
    ws['E2'].alignment = Alignment(vertical='top')
    for nm, f in (('Epargne', 'SUM(V_Montant)'), ('Especes', 'MAX(0,SUM(V_Montant)-SUM(A_Debit)+SUM(A_Remb))'), ('Valorisation', 'TotalTitres+Especes')):
        name(wb, nm, f)
    cards = [  # (titre, formule, format, libellé secondaire, formule secondaire, format secondaire)
        ('Valorisation totale', '=Valorisation', EUR, 'dont espèces', '=Especes', EUR),
        ('Plus-value latente', '=Valorisation-Epargne', EUR_S, 'rendement global', '=IF(Epargne>0,(Valorisation-Epargne)/Epargne,0)', PCT_S),
        ('Rendement annualisé (TRI)', '=IFERROR(XIRR(TRI_Flux,TRI_Dates),"")', PCT_S, 'virements', '=COUNT(V_Date)', INT),
        ('Montant épargné', '=Epargne', EUR, 'investi en titres', '=SUM(A_Brut)', EUR),
        ('Parts détenues', f'=Portefeuille!D{pt}', INT, 'lignes', f'=COUNTIF(Portefeuille!D{p0}:D{pt - 1},">0")', INT),
        ('Frais nets', '=SUM(A_Frais)-SUM(A_Remb)', EUR, 'remboursés', '=SUM(A_Remb)', EUR),
        ('Frais moyens par achat', '=IFERROR(AVERAGE(A_FraisPct),0)', PCT2, 'part remboursée', '=IF(SUM(A_Frais)>0,SUM(A_Remb)/SUM(A_Frais),0)', '0%'),
        ('Solde espèces', '=Especes', EUR, 'achats enregistrés', '=COUNT(A_Date)', INT),
    ]
    # 4 cartes par ligne, 2 colonnes chacune (B:C, D:E, F:G, H:I), séparées par un filet blanc épais
    ws.column_dimensions['A'].width = 2
    for c in 'BCDEFGHI':
        ws.column_dimensions[c].width = 13.5
    gap = Side(style='thick', color='FFFFFF')
    starts = ['B', 'D', 'F', 'H']
    for k, (lab, f, fmt, sublab, subf, subfmt) in enumerate(cards):
        row = 5 if k < 4 else 10
        c1 = starts[k % 4]
        c2 = L(ws[c1 + '1'].column + 1)
        hero = k == 0
        bg, fg, sfg = (PETROL, 'FFFFFF', 'BFD3D0') if hero else (BG, TEXT, MODERATE)
        for rr in range(row, row + 4):
            ws[f'{c1}{rr}'].fill = fill(bg)
            ws[f'{c2}{rr}'].fill = fill(bg)
            ws[f'{c1}{rr}'].border = Border(left=gap)
            ws[f'{c2}{rr}'].border = Border(right=gap)
        ws.merge_cells(f'{c1}{row}:{c2}{row}')
        ws.merge_cells(f'{c1}{row + 1}:{c2}{row + 2}')
        ws[f'{c1}{row}'] = lab.upper()
        ws[f'{c1}{row}'].font = Font(name=FONT, size=8, bold=True, color=sfg)
        ws[f'{c1}{row}'].alignment = Alignment(indent=1, vertical='bottom')
        ws[f'{c1}{row + 1}'] = f
        ws[f'{c1}{row + 1}'].number_format = fmt
        ws[f'{c1}{row + 1}'].font = Font(name=FONT, size=18, bold=True, color=fg)
        ws[f'{c1}{row + 1}'].alignment = Alignment(indent=1, vertical='center')
        ws[f'{c1}{row + 3}'] = sublab
        ws[f'{c1}{row + 3}'].font = Font(name=FONT, size=8, color=sfg)
        ws[f'{c1}{row + 3}'].alignment = Alignment(indent=1, vertical='top')
        ws[f'{c2}{row + 3}'] = subf
        ws[f'{c2}{row + 3}'].number_format = subfmt
        ws[f'{c2}{row + 3}'].font = Font(name=FONT, size=8, bold=True, color=fg)
        ws[f'{c2}{row + 3}'].alignment = Alignment(horizontal='right', vertical='top', indent=1)
    for rr, h in ((5, 18), (6, 18), (7, 18), (8, 18), (9, 8), (10, 18), (11, 18), (12, 18), (13, 18)):
        ws.row_dimensions[rr].height = h
    ws.conditional_formatting.add('D6', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, size=18, bold=True, color=CRIT)))
    ws.conditional_formatting.add('D6', CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=FONT, size=18, bold=True, color=SUCCESS)))
    # tableau réel / cible
    dash.section('B15', 'Allocation réelle et cible')
    dash.header(16, 2, ['Indice', 'Ticker', 'Valeur', 'Réel', 'Cible', 'Écart', 'À acheter'], height=22)
    for i, t in enumerate(order):
        r = 17 + i
        pr = p0 + i
        dash.put(r, 2, f'=Portefeuille!A{pr}', bold=True)
        dash.put(r, 3, t, align='center')
        ws.cell(r, 3).fill = fill(colors[t])
        ws.cell(r, 3).font = font(10, True, 'FFFFFF' if colors[t] not in ('F8D10A', 'A0D7A2', 'BAE3EC') else TEXT)
        dash.put(r, 4, f'=Portefeuille!H{pr}', EUR)
        dash.put(r, 5, f'=Portefeuille!K{pr}', PCT1, align='center')
        dash.put(r, 6, f'=Portefeuille!L{pr}', PCT1, align='center')
        dash.put(r, 7, f'=Portefeuille!M{pr}', PTS, align='center')
        dash.put(r, 8, f'=Portefeuille!P{pr}', INT, align='center', bold=True, color=PRIMARY)
    tl = 16 + len(order)
    ws.conditional_formatting.add(f'G17:G{tl}', CellIsRule(operator='lessThan', formula=['-5'], font=Font(name=FONT, color=CRIT, bold=True)))
    ws.conditional_formatting.add(f'G17:G{tl}', CellIsRule(operator='greaterThan', formula=['5'], font=Font(name=FONT, color='8A4B00', bold=True)))
    # anneau : répartition réelle
    dn = DoughnutChart(holeSize=58)
    style_chart(dn, 'Répartition réelle', legend='r')
    dn.add_data(Reference(ws, min_col=5, min_row=17, max_row=tl), titles_from_data=False)
    dn.set_categories(Reference(ws, min_col=3, min_row=17, max_row=tl))
    for i, t in enumerate(order):
        pt_ = DataPoint(idx=i)
        pt_.graphicalProperties.solidFill = colors[t]
        pt_.graphicalProperties.line.solidFill = 'FFFFFF'
        dn.series[0].dPt.append(pt_)
    dn.width, dn.height = 10, 7.5
    ws.add_chart(dn, f'B{tl + 2}')
    bc = BarChart()
    style_chart(bc, 'Réel et cible par ETF')
    bc.add_data(Reference(ws, min_col=5, min_row=16, max_row=tl), titles_from_data=True)
    bc.add_data(Reference(ws, min_col=6, min_row=16, max_row=tl), titles_from_data=True)
    color_series(bc.series[0], PRIMARY)
    color_series(bc.series[1], DIVIDER)
    bc.set_categories(Reference(ws, min_col=3, min_row=17, max_row=tl))
    bc.y_axis.number_format = '0%'
    bc.gapWidth = 80
    bc.width, bc.height = 13, 7.5
    ws.add_chart(bc, f'F{tl + 2}')
    # courbe épargne / valeur
    lc2 = LineChart()
    style_chart(lc2, 'Montant épargné et valeur réelle')
    lc2.add_data(Reference(ev.ws, min_col=4, min_row=4, max_row=m_last), titles_from_data=True)
    color_series(lc2.series[0], MODERATE, line=True, dash='dash', width=19050)
    lc2.add_data(Reference(ev.ws, min_col=cV, min_row=4, max_row=m_last), titles_from_data=True)
    color_series(lc2.series[1], PRIMARY, line=True, width=28575)
    lc2.set_categories(Reference(ev.ws, min_col=1, min_row=5, max_row=m_last))
    lc2.x_axis.number_format = 'mmm yy'
    lc2.y_axis.number_format = '#,##0 "€"'
    lc2.width, lc2.height = 23.5, 8
    ws.add_chart(lc2, f'B{tl + 18}')
    leg = tl + 35
    ws[f'B{leg}'] = 'Mode d’emploi'
    ws[f'B{leg}'].font = font(9, True, NEUTRAL)
    notes = ['Les cellules en bleu se saisissent (Historique, Mes ETF, Cours, Compositions, Simulateur) ; tout le reste se calcule.',
             'Valorisation = valeur des titres au dernier cours + solde espèces. Solde espèces = virements − débits courtier + remboursements de frais.',
             'TRI : taux de rendement interne annualisé sur les virements (XIRR, base 365 jours).',
             'Fichier généré par PEAnalyze : les données de l’application font foi.']
    for i, n in enumerate(notes):
        ws[f'B{leg + 1 + i}'] = '• ' + n
        ws[f'B{leg + 1 + i}'].font = font(8, color=MODERATE)
    ws.sheet_view.zoomScale = 100

    # mise en page impression
    for s in wb.worksheets:
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.print_options.horizontalCentered = True
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)
    print('écrit', dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
